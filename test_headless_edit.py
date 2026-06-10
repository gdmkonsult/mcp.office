from pathlib import Path
import asyncio
from docx import Document
import server


def run_docx_test() -> None:
    root = Path("test-files")
    root.mkdir(exist_ok=True)
    source = root / "headless-source.docx"
    out = root / "headless-edited.docx"

    doc = Document()
    doc.add_paragraph("Project Eneo draft contract")
    doc.add_paragraph("Old company name: Acme AB")
    doc.add_paragraph("Delete me paragraph marker")
    doc.save(source)

    prompt = "\n".join(
        [
            'replace "Acme AB" with "Eneo AI Services AB"',
            'append paragraph "Approved by assistant workflow"',
            'delete paragraph containing "Delete me paragraph marker"',
        ]
    )

    result = asyncio.run(server.edit_document_from_prompt(str(source), str(out), prompt))
    print("DOCX_RESULT_OK", result.get("ok"))
    print("DOCX_SUMMARY", result.get("summary"))

    edited = Document(out)
    texts = [p.text for p in edited.paragraphs]
    print("DOCX_HAS_NEW_NAME", any("Eneo AI Services AB" in t for t in texts))
    print("DOCX_HAS_APPEND", any("Approved by assistant workflow" in t for t in texts))
    print("DOCX_HAS_DELETE_MARKER", any("Delete me paragraph marker" in t for t in texts))


def run_pdf_test() -> None:
    root = Path("test-files")
    source = root / "existing-sample.pdf"
    out = root / "headless-edited.pdf"

    prompt = "ClientName=Eneo\nContractDate=2026-06-10"
    result = asyncio.run(server.edit_document_from_prompt(str(source), str(out), prompt))
    print("PDF_RESULT_OK", result.get("ok"))
    print("PDF_UPDATED_FIELDS", result.get("updated_fields"))
    if not result.get("ok"):
        print("PDF_ERROR", result.get("error"))


if __name__ == "__main__":
    run_docx_test()
    run_pdf_test()
