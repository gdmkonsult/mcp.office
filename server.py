import contextlib
import os
import re
from pathlib import Path
from typing import Any

from mcp.server.fastmcp import FastMCP
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from starlette.applications import Starlette
from starlette.responses import JSONResponse
from starlette.routing import Mount, Route
from docx import Document

mcp = FastMCP(
    "Office",
    host="0.0.0.0",
    stateless_http=True,
    json_response=True,
)


def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def _parse_prompt_actions(prompt: str) -> list[dict[str, str]]:
    actions: list[dict[str, str]] = []

    patterns = [
        (r"replace\s+\"(?P<old>.+?)\"\s+with\s+\"(?P<new>.+?)\"", "replace"),
        (r"append\s+paragraph\s+\"(?P<text>.+?)\"", "append_paragraph"),
        (r"prepend\s+paragraph\s+\"(?P<text>.+?)\"", "prepend_paragraph"),
        (r"delete\s+paragraph\s+containing\s+\"(?P<contains>.+?)\"", "delete_paragraph"),
    ]

    for line in [item.strip() for item in prompt.splitlines() if item.strip()]:
        matched = False
        for pattern, kind in patterns:
            m = re.search(pattern, line, flags=re.IGNORECASE)
            if m:
                action = {"type": kind}
                action.update({k: v for k, v in m.groupdict().items() if v is not None})
                actions.append(action)
                matched = True
                break
        if not matched and "=" in line:
            # Allow PDF form filling with "Field Name=Value" lines.
            name, value = line.split("=", 1)
            if name.strip():
                actions.append(
                    {
                        "type": "set_form_field",
                        "field": name.strip(),
                        "value": value.strip(),
                    }
                )

    return actions


def _edit_docx_from_actions(input_path: Path, output_path: Path, actions: list[dict[str, str]]) -> dict[str, Any]:
    doc = Document(str(input_path))
    summary: list[str] = []

    for action in actions:
        kind = action.get("type")
        if kind == "replace":
            old = action.get("old", "")
            new = action.get("new", "")
            replaced = 0
            for paragraph in doc.paragraphs:
                if old and old in paragraph.text:
                    paragraph.text = paragraph.text.replace(old, new)
                    replaced += 1
            summary.append(f"replace:{old}->{new} paragraphs:{replaced}")

        elif kind == "append_paragraph":
            text = action.get("text", "")
            doc.add_paragraph(text)
            summary.append("append_paragraph")

        elif kind == "prepend_paragraph":
            text = action.get("text", "")
            first = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph("")
            inserted = first.insert_paragraph_before(text)
            inserted.style = first.style
            summary.append("prepend_paragraph")

        elif kind == "delete_paragraph":
            contains = action.get("contains", "")
            deleted = 0
            for paragraph in list(doc.paragraphs):
                if contains and contains in paragraph.text:
                    element = paragraph._element
                    parent = element.getparent()
                    if parent is not None:
                        parent.remove(element)
                        deleted += 1
            summary.append(f"delete_paragraph:{contains} count:{deleted}")

    _ensure_parent(output_path)
    doc.save(str(output_path))

    return {
        "ok": True,
        "output_path": str(output_path),
        "summary": summary,
    }


def _edit_pdf_form_from_actions(input_path: Path, output_path: Path, actions: list[dict[str, str]]) -> dict[str, Any]:
    field_updates = {
        action["field"]: action["value"]
        for action in actions
        if action.get("type") == "set_form_field"
    }
    if not field_updates:
        return {
            "ok": False,
            "error": "No PDF form field actions found. Use lines like FieldName=Value in prompt.",
        }

    reader = PdfReader(str(input_path))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)

    try:
        writer.update_page_form_field_values(writer.pages[0], field_updates)
    except Exception as exc:
        return {
            "ok": False,
            "error": f"Failed to apply PDF form values: {exc}",
        }

    _ensure_parent(output_path)
    with open(output_path, "wb") as f:
        writer.write(f)

    return {
        "ok": True,
        "output_path": str(output_path),
        "updated_fields": list(field_updates.keys()),
    }


def _rewrite_pdf_text_from_actions(input_path: Path, output_path: Path, actions: list[dict[str, str]]) -> dict[str, Any]:
    reader = PdfReader(str(input_path))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    lines = [line for line in text.splitlines() if line.strip()]

    summary: list[str] = []
    for action in actions:
        kind = action.get("type")
        if kind == "replace":
            old = action.get("old", "")
            new = action.get("new", "")
            lines = [line.replace(old, new) for line in lines]
            summary.append(f"replace:{old}->{new}")
        elif kind == "delete_paragraph":
            contains = action.get("contains", "")
            before = len(lines)
            lines = [line for line in lines if contains not in line]
            summary.append(f"delete_lines:{before - len(lines)}")
        elif kind in {"append_paragraph", "prepend_paragraph"}:
            text_value = action.get("text", "")
            if kind == "append_paragraph":
                lines.append(text_value)
            else:
                lines.insert(0, text_value)
            summary.append(kind)

    _ensure_parent(output_path)
    c = canvas.Canvas(str(output_path), pagesize=A4)
    width, height = A4
    y = height - 50
    for line in lines:
        if y < 50:
            c.showPage()
            y = height - 50
        c.drawString(40, y, line[:150])
        y -= 16
    c.save()

    return {
        "ok": True,
        "output_path": str(output_path),
        "summary": summary,
        "warning": "Original PDF layout is not preserved in text-rewrite fallback mode.",
    }


@mcp.tool()
async def describe_capabilities() -> dict:
    """Describe what this MCP supports for document editing."""
    return {
        "server": "mcp.office",
        "capabilities": {
            "edit_docx": "Headless editing of Word documents (.docx) using python-docx",
            "edit_pdf_forms": "PDF form field filling using PyPDF",
            "edit_pdf_text": "PDF text rewriting using reportlab (fallback for non-form PDFs)",
        },
        "tools": ["edit_document_from_prompt"],
        "headless_only": True,
        "backend_independent": True,
    }


@mcp.tool()
async def edit_document_from_prompt(input_path: str, output_path: str, prompt: str) -> dict:
    """Headless edit of DOCX or PDF (forms) based on a natural-language prompt.

    Prompt syntax currently supports:
    - replace "old" with "new"
    - append paragraph "text"
    - prepend paragraph "text"
    - delete paragraph containing "text"
    - PDF form filling with lines like: FieldName=Value
    """
    src = Path(input_path)
    dst = Path(output_path)

    if not src.exists():
        return {"ok": False, "error": f"Input file does not exist: {input_path}"}

    actions = _parse_prompt_actions(prompt)
    if not actions:
        return {
            "ok": False,
            "error": "No supported actions parsed from prompt.",
            "hint": "For DOCX: replace/append/prepend/delete paragraph commands. For PDF forms: Field=Value lines.",
        }

    suffix = src.suffix.lower()
    if suffix == ".docx":
        result = _edit_docx_from_actions(src, dst, actions)
        result["document_type"] = "docx"
        result["actions"] = actions
        return result

    if suffix == ".pdf":
        has_form_actions = any(action.get("type") == "set_form_field" for action in actions)
        if has_form_actions:
            result = _edit_pdf_form_from_actions(src, dst, actions)
        else:
            result = _rewrite_pdf_text_from_actions(src, dst, actions)

        if not result.get("ok") and "No /AcroForm" in str(result.get("error", "")):
            result = _rewrite_pdf_text_from_actions(src, dst, actions)
        result["document_type"] = "pdf"
        result["actions"] = actions
        return result

    return {
        "ok": False,
        "error": f"Unsupported extension: {suffix}",
        "supported": [".docx", ".pdf"],
    }


async def healthz(_request):
    return JSONResponse({"status": "ok"})


async def readyz(_request):
    return JSONResponse({"status": "ready"})


@contextlib.asynccontextmanager
async def lifespan(app: Starlette):
    async with mcp.session_manager.run():
        yield


app = Starlette(
    routes=[
        Route("/healthz", healthz),
        Route("/readyz", readyz),
        Mount("/", app=mcp.streamable_http_app()),
    ],
    lifespan=lifespan,
)

if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
